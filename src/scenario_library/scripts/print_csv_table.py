import pandas as pd
from openpyxl.styles import Alignment, Border, Side

CSV_FILE = "/home/hzq/桌面/测试结果/reckless_driving_identification_test_log_2025-12-15_09-40-13.csv"
EXCEL_FILE = "/home/hzq/桌面/测试结果/scenario_results_table.xlsx"

def csv_to_excel(csv_path, excel_path):
    # 读取 CSV
    df = pd.read_csv(csv_path)

    with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Results")

        ws = writer.book["Results"]

        # ================================
        # 样式定义
        # ================================
        center_align = Alignment(horizontal="center", vertical="center")

        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        # ================================
        # 设置居中 + 全边框
        # ================================
        for row in ws.iter_rows(
            min_row=1,
            max_row=ws.max_row,
            min_col=1,
            max_col=ws.max_column
        ):
            for cell in row:
                cell.alignment = center_align
                cell.border = thin_border

        # ================================
        # 自动调整列宽
        # ================================
        for col in ws.columns:
            max_length = 0
            col_letter = col[0].column_letter
            for cell in col:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = max_length + 2

    print(f"✅ 表格文件已生成（居中 + 全边框）：{excel_path}")


if __name__ == "__main__":
    csv_to_excel(CSV_FILE, EXCEL_FILE)
